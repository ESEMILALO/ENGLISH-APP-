"""Put the phrasal verbs read out of the book into the spreadsheet.

`read_book.py` leaves `tools/book_verbs.json`: the verb, the meaning and
the example exactly as the book prints them. That is not yet a row in
this spreadsheet, which wants the Spanish too, a pronunciation a Spanish
reader can say, and the verb's forms in the notes. So each one goes to
Claude in the same style as every other row, with the book's own meaning
and example as the thing being dressed, not replaced.

A verb already somewhere in the workbook is left alone and reported, so
a book that overlaps the lists you already have does not fill the sheet
with second copies.

    python tools/add_book_words.py                 # into School vocabulary
    python tools/add_book_words.py --sheet Series
    python tools/add_book_words.py --dry-run       # say what it would do
"""

import argparse
import datetime
import json
import shutil
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "server"))

import glossary  # noqa: E402
from glossary import api_key, normalise, tidy  # noqa: E402

VERBS = HERE / "book_verbs.json"
READY = HERE / "book_rows.json"
API_URL = "https://api.anthropic.com/v1/messages"
MODEL = "claude-sonnet-5"
API_VERSION = "2023-06-01"
BATCH = 8

PROMPT = """You are filling in rows of a Spanish speaker's English vocabulary \
spreadsheet. Match the existing style exactly.

Two real rows, so you can see the style:

Word: Supper
Pronunciation: super
Meaning 1: Evening meal — Cena.
Example: We had soup for supper.
Notes: More common in some regions than dinner.

Word: Jot down
Pronunciation: yat daun
Meaning 1: Write quickly — Anotar.
Example: Jot down the phone number.
Notes: Regular verb. Base: jot down — Present: jots down — Past: jotted \
down — Past participle: jotted down. Separable: jot something down.

Each entry you are given comes from a phrasal verb book and already has a \
meaning and an example. Keep them. Your job is to put them into this \
sheet's shape, not to think of better ones.

Rules:
- "word": the phrasal verb, capitalised the way the sheet does it -- first \
letter only, like "Wake up" or "Jot down". Never "Wake Up", never "to wake up".
- "pronunciation": respelled so a Spanish reader says it correctly. Not IPA. \
Spanish spelling conventions, with an accent marking the stress where Spanish \
would use one.
- "meaning": the book's meaning, trimmed to a short gloss, then the em dash, \
then the Spanish for it, then a full stop. "to stop sleeping" becomes \
"Stop sleeping — Despertarse." Keep the book's sense exactly; do not \
substitute another sense of the verb.
- "example": the book's example sentence, exactly as given, with a full stop \
if it is missing one. Do not write a new one.
- "notes": the verb's forms in the style above, plus whether it is separable \
and anything genuinely worth knowing. Keep it to one or two sentences.

Return ONLY a JSON array, one object per entry, in the order given:
[{"n": <the number given>, "word": "...", "pronunciation": "...", \
"meaning": "...", "example": "...", "notes": "..."}]"""


def first_json(text):
    if text.startswith("```"):
        text = text.split("\n", 1)[1].rsplit("```", 1)[0]
    at = text.find("[")
    while at >= 0:
        depth, in_str, esc = 0, False, False
        for i in range(at, len(text)):
            ch = text[i]
            if in_str:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == '"':
                    in_str = False
                continue
            if ch == '"':
                in_str = True
            elif ch == "[":
                depth += 1
            elif ch == "]":
                depth -= 1
                if depth == 0:
                    try:
                        return json.loads(text[at:i + 1])
                    except Exception:
                        break
        at = text.find("[", at + 1)
    return None


def ask(batch, key, attempts=4):
    lines = []
    for n, item in enumerate(batch, 1):
        lines.append("%d. %s\n   meaning: %s\n   example: %s"
                     % (n, item["verb"], item["meaning"], item["example"]))
    body = json.dumps({
        "model": MODEL,
        "max_tokens": 4000,
        "system": PROMPT,
        "messages": [{"role": "user", "content": "\n".join(lines)}],
    }).encode("utf-8")

    last = None
    for attempt in range(attempts):
        req = urllib.request.Request(API_URL, data=body, method="POST", headers={
            "content-type": "application/json",
            "x-api-key": key,
            "anthropic-version": API_VERSION,
        })
        try:
            with urllib.request.urlopen(req, timeout=120) as res:
                payload = json.loads(res.read().decode("utf-8"))
            text = "".join(p.get("text", "") for p in payload.get("content", []))
            got = first_json(text.strip())
            if got is None:
                raise ValueError("no rows in the reply")
            return got
        except urllib.error.HTTPError as e:
            # A 400 carries the reason with it, and losing that reason
            # means losing eight words to a mystery.
            try:
                detail = e.read().decode("utf-8")[:300]
            except Exception:
                detail = ""
            last = RuntimeError("HTTP %s %s" % (e.code, detail)) if detail else e
            if e.code not in (408, 409, 429, 500, 502, 503, 504, 529):
                raise last
        except Exception as e:
            last = e
        time.sleep(3 * (attempt + 1))
    raise last


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--sheet", default="School vocabulary")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not VERBS.exists():
        sys.exit("Run  python tools/read_book.py  first.")
    book = json.loads(VERBS.read_text(encoding="utf-8"))
    print("%d entries read from the book" % len(book))

    # the same verb can be taught twice in one book
    seen, unique = set(), []
    for row in book:
        key = normalise(row["verb"])
        if key and key not in seen:
            seen.add(key)
            unique.append(row)
    print("%d different verbs" % len(unique))

    # anything already anywhere in the workbook stays where it is
    import openpyxl
    wb = openpyxl.load_workbook(glossary.WORKBOOK, read_only=True)
    have = {}
    for name in wb.sheetnames:
        for cells in wb[name].iter_rows(min_row=2, max_col=1, values_only=True):
            if cells[0]:
                have.setdefault(normalise(cells[0]), name)
    wb.close()

    fresh, already = [], []
    for row in unique:
        where = have.get(normalise(row["verb"]))
        (already if where else fresh).append((row, where))
    print("%d already in the sheets, %d new" % (len(already), len(fresh)))
    if already[:6]:
        for row, where in already[:6]:
            print("   have: %-18s (%s)" % (row["verb"], where))
        if len(already) > 6:
            print("   ... and %d more" % (len(already) - 6))

    todo = [r for r, _ in fresh]
    if args.limit:
        todo = todo[:args.limit]
    if not todo:
        print("\nNothing new to add.")
        return

    key = api_key()
    if not key:
        sys.exit("This needs your Anthropic key -- the same one the app uses.")

    # Rows already prepared on an earlier run are not paid for twice.
    rows = []
    if READY.exists():
        try:
            rows = json.loads(READY.read_text(encoding="utf-8"))
            print("%d rows were already prepared" % len(rows))
        except Exception:
            rows = []
    done_words = {normalise(r["word"]) for r in rows}
    todo = [t for t in todo if normalise(t["verb"]) not in done_words]

    for start in range(0, len(todo), BATCH):
        batch = todo[start:start + BATCH]
        try:
            got = ask(batch, key)
        except Exception as e:
            print("  a batch failed (%s) -- carrying on" % str(e)[:60], flush=True)
            continue
        for item in got:
            word = str(item.get("word") or "").strip()
            meaning = tidy(str(item.get("meaning") or "").strip())
            if not word or not meaning:
                continue
            rows.append({
                "word": word,
                "pronunciation": str(item.get("pronunciation") or "").strip(),
                "meaning": meaning,
                "example": str(item.get("example") or "").strip(),
                "notes": tidy(str(item.get("notes") or "").strip()),
            })
        # Written as they are made, so a stumble at the end -- or a
        # closed laptop -- does not throw away an hour of asking.
        READY.write_text(json.dumps(rows, indent=1, ensure_ascii=False),
                         encoding="utf-8")
        print("  prepared %d/%d" % (min(start + BATCH, len(todo)), len(todo)),
              flush=True)

    print("\n%d rows ready. First few:" % len(rows))
    for r in rows[:5]:
        print("   %-16s %-14s %s" % (r["word"], r["pronunciation"], r["meaning"][:52]))

    if args.dry_run:
        out = HERE / "book_rows_preview.json"
        out.write_text(json.dumps(rows, indent=1, ensure_ascii=False), encoding="utf-8")
        print("\nDry run. Nothing written. The rows are in %s" % out.name)
        return

    # One copy of the workbook before any of it, and one save after all of
    # it -- three hundred saves would be three hundred chances to corrupt.
    glossary.BACKUPS.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    kept = glossary.BACKUPS / ("ENGLISH SCHOOL.pre-book-%s.xlsx" % stamp)
    shutil.copy(glossary.WORKBOOK, kept)
    print("\nWorkbook copied to backups/%s" % kept.name)

    wb = openpyxl.load_workbook(glossary.WORKBOOK)
    if args.sheet not in wb.sheetnames:
        sys.exit("No sheet called %r. There is: %s" % (args.sheet, ", ".join(wb.sheetnames)))
    ws = wb[args.sheet]
    at = ws.max_row
    for r in rows:
        at += 1
        ws.cell(row=at, column=1).value = r["word"]
        ws.cell(row=at, column=2).value = r["pronunciation"] or None
        ws.cell(row=at, column=3).value = r["meaning"]
        ws.cell(row=at, column=4).value = r["example"] or None
        ws.cell(row=at, column=glossary.NOTES_COL).value = r["notes"] or None
    wb.save(glossary.WORKBOOK)

    print("Added %d verbs to %s (now %d rows)." % (len(rows), args.sheet, at))
    # They are in the workbook now; keeping the working copy would make a
    # second run add them all over again.
    try:
        READY.unlink()
    except OSError:
        pass
    print("Run  python tools/build_word_log.py  to put them in the app.")


if __name__ == "__main__":
    main()
