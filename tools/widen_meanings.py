"""Give every sheet five Meaning slots instead of three.

The Ultimate Phrasal Verb Book gives some verbs far more senses than
three -- pick up has twelve -- and three slots meant throwing away 123
of them. Two more pairs of columns are inserted before Notes, so the
sheet reads:

    Word | Pron | M1 E1 | M2 E2 | M3 E3 | M4 E4 | M5 E5 | Notes

Notes moves from column 9 to column 13. Nothing else moves: no word
changes sheet or spelling, so no id changes and no progress is lost.

    python tools/widen_meanings.py --dry-run
    python tools/widen_meanings.py
"""

import argparse
import datetime
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "server"))
import glossary  # noqa: E402

OLD_NOTES, NEW_NOTES = 9, 13
HEADER = ["Word/expression", "Pronunciation",
          "Meaning 1", "Example", "Meaning 2", "Example", "Meaning 3", "Example",
          "Meaning 4", "Example", "Meaning 5", "Example", "Notes"]


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(glossary.WORKBOOK)

    first = wb[wb.sheetnames[0]]
    if first.cell(row=1, column=NEW_NOTES).value == "Notes":
        sys.exit("The sheets are already five meanings wide.")

    moved = 0
    for name in wb.sheetnames:
        ws = wb[name]
        for row in range(1, ws.max_row + 1):
            note = ws.cell(row=row, column=OLD_NOTES).value
            if note is not None:
                ws.cell(row=row, column=NEW_NOTES).value = note
                ws.cell(row=row, column=OLD_NOTES).value = None
                moved += 1
        for i, title in enumerate(HEADER, 1):
            ws.cell(row=1, column=i).value = title
        print("  %-24s %d rows" % (name, ws.max_row - 1))

    print("%d notes moved from column %d to column %d" % (moved, OLD_NOTES, NEW_NOTES))
    if args.dry_run:
        print("\nDry run. Nothing written.")
        return

    glossary.BACKUPS.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    kept = glossary.BACKUPS / ("ENGLISH SCHOOL.pre-widen-%s.xlsx" % stamp)
    shutil.copy(glossary.WORKBOOK, kept)
    wb.save(glossary.WORKBOOK)
    print("Saved. Copy before this: backups/%s" % kept.name)


if __name__ == "__main__":
    main()
