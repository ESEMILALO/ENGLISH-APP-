"""Put "504 Main words" and "504 Secondary" into one topic.

The book teaches 504 words and, beside each one, the other ways of saying
it. Those arrived here as two separate lists, which meant the app treated
"abandon" and "desert" as belonging to different subjects. They are the
same lesson, so they become one sheet.

A word's id in the app is "<sheet name>::<word>", so renaming a sheet
would quietly orphan everything already learnt from it. The ids are
migrated in the saved progress file here, and the app does the same for
the copy in the browser the next time it starts.

    python tools/merge_504_sheets.py --dry-run
    python tools/merge_504_sheets.py
"""

import argparse
import datetime
import json
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "server"))

import glossary  # noqa: E402
from glossary import normalise  # noqa: E402

OLD = ["504 Main words", "504 Secondary"]
NEW = "504 Essential words"
PROGRESS = ROOT / "progress" / "progress.json"


def migrate(value, pairs):
    """Rewrite every old id anywhere inside a saved structure.

    Ids turn up as dictionary keys (progress, speech scores), as strings
    in lists (today's set, my list, the words a conversation used) and
    singly (the phrasal verb of the day). Walking the whole thing is
    shorter than naming each place, and does not go stale when a new
    place starts holding an id.
    """
    if isinstance(value, str):
        for old, new in pairs:
            if value.startswith(old):
                return new + value[len(old):]
        return value
    if isinstance(value, list):
        return [migrate(v, pairs) for v in value]
    if isinstance(value, dict):
        return {migrate(k, pairs): migrate(v, pairs) for k, v in value.items()}
    return value


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(glossary.WORKBOOK)
    for name in OLD:
        if name not in wb.sheetnames:
            sys.exit("No sheet called %r. There is: %s"
                     % (name, ", ".join(wb.sheetnames)))
    if NEW in wb.sheetnames:
        sys.exit("%r already exists -- this has been run before." % NEW)

    header = [c.value for c in wb[OLD[0]][1]]
    rows, seen, dropped = [], set(), []
    for name in OLD:
        for row in wb[name].iter_rows(min_row=2, values_only=True):
            if not row or not row[0] or not str(row[0]).strip():
                continue
            key = normalise(row[0])
            if key in seen:
                dropped.append((name, row[0]))
                continue
            seen.add(key)
            rows.append(list(row))

    print("%s: %d rows" % (OLD[0], sum(1 for r in wb[OLD[0]].iter_rows(min_row=2, values_only=True) if r and r[0])))
    print("%s: %d rows" % (OLD[1], sum(1 for r in wb[OLD[1]].iter_rows(min_row=2, values_only=True) if r and r[0])))
    print("combined: %d rows (%d second copies left out: %s)"
          % (len(rows), len(dropped), ", ".join(w for _, w in dropped) or "none"))

    # what this does to the ids already saved
    pairs = [(o + "::", NEW + "::") for o in OLD]
    if PROGRESS.exists():
        saved = json.loads(PROGRESS.read_text(encoding="utf-8"))
        before = json.dumps(saved, ensure_ascii=False)
        after = json.dumps(migrate(saved, pairs), ensure_ascii=False)
        moved = sum(1 for k in (saved.get("data", {}).get("progress") or {})
                    if any(k.startswith(o) for o, _ in pairs))
        print("saved progress: %d words carried over" % moved)
    else:
        saved, after = None, None
        print("saved progress: no file to carry over")

    if args.dry_run:
        print("\nDry run. Nothing written.")
        return

    glossary.BACKUPS.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    kept = glossary.BACKUPS / ("ENGLISH SCHOOL.pre-merge-%s.xlsx" % stamp)
    shutil.copy(glossary.WORKBOOK, kept)
    print("\nWorkbook copied to backups/%s" % kept.name)

    # The new sheet takes the first one's place in the order, so the app's
    # topic list does not suddenly rearrange itself.
    at = wb.sheetnames.index(OLD[0])
    ws = wb.create_sheet(NEW, at)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for name in OLD:
        del wb[name]
    wb.save(glossary.WORKBOOK)
    print("Wrote %r with %d rows; removed %s." % (NEW, len(rows), " and ".join(repr(o) for o in OLD)))

    if saved is not None:
        copy = PROGRESS.with_name("progress.pre-merge-%s.json" % stamp)
        shutil.copy(PROGRESS, copy)
        PROGRESS.write_text(after, encoding="utf-8")
        print("Progress ids migrated (old file kept as %s)." % copy.name)

    print("\nRun  python tools/build_word_log.py  to put it in the app.")


if __name__ == "__main__":
    main()
