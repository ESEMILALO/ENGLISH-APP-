"""
Removes duplicate entries for the same word.

Two kinds exist: the same word twice in one sheet, and the same word in
several sheets. Both mean studying identical material as separate cards,
and -- because progress is keyed per category -- mastering it two or three
times over.

For each group the richest row wins: the one with the most filled cells,
then the most text. Ties fall back to sheet order, so the earlier list
keeps the word.

It is a merge, not a delete: a duplicate often carries a sense the winner
lacks ("wind up" as "turn a mechanism"), so distinct meanings are copied
into the winner's free slots before the other rows go. A word with more
distinct senses than the three available slots is left completely alone
rather than silently truncated. Everything is written to
dedupe_report.txt.

Run without --apply for a dry run.
"""
import os
import re
import shutil
import sys
from collections import defaultdict
from datetime import datetime

import openpyxl

MEANING_COLS = [(3, 4), (5, 6), (7, 8)]  # (meaning, example) column pairs


def meaning_key(text):
    """Compares the English half only, ignoring punctuation and case."""
    english = text.split("—")[0]
    return re.sub(r"[^a-z ]", "", english.lower()).strip()

FILE = "ENGLISH SCHOOL.xlsx"
SHEET_ORDER = ["School vocabulary", "504 Main words", "504 Secondary", "School books", "Series"]


def richness(ws, row):
    filled = 0
    length = 0
    for col in range(2, 10):  # pronunciation, meanings, examples, notes
        v = ws.cell(row=row, column=col).value
        if v is not None and str(v).strip():
            filled += 1
            length += len(str(v).strip())
    return filled, length


def main(apply_changes):
    if apply_changes:
        os.makedirs("backups", exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        shutil.copy(FILE, os.path.join("backups", "ENGLISH SCHOOL.pre-dedupe-%s.xlsx" % stamp))

    wb = openpyxl.load_workbook(FILE)
    groups = defaultdict(list)

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in range(2, ws.max_row + 1):
            raw = ws.cell(row=row, column=1).value
            if raw is None or not str(raw).strip():
                continue
            word = str(raw).strip()
            filled, length = richness(ws, row)
            order = SHEET_ORDER.index(sheet) if sheet in SHEET_ORDER else 99
            groups[word.lower()].append(
                {"sheet": sheet, "row": row, "word": word,
                 "filled": filled, "length": length, "order": order}
            )

    report = []
    to_delete = defaultdict(list)
    merges = []            # (sheet, row, [(meaning, example), ...]) to write into the winner
    same_sheet = cross_sheet = skipped = moved = 0

    for word, entries in sorted(groups.items()):
        if len(entries) < 2:
            continue
        # richest first; ties break toward the earlier sheet, then earlier row
        ranked = sorted(entries, key=lambda e: (-e["filled"], -e["length"], e["order"], e["row"]))
        keep, drop = ranked[0], ranked[1:]

        # gather every distinct sense across the group, winner's first
        senses = []
        seen_keys = set()
        for entry in ranked:
            ws = wb[entry["sheet"]]
            for mc, ec in MEANING_COLS:
                m = ws.cell(row=entry["row"], column=mc).value
                if not m or not str(m).strip():
                    continue
                key = meaning_key(str(m).strip())
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                e = ws.cell(row=entry["row"], column=ec).value
                senses.append((str(m).strip(), str(e).strip() if e else ""))

        if len(senses) > len(MEANING_COLS):
            skipped += 1
            report.append("%s -- LEFT ALONE: %d distinct senses, only %d slots"
                          % (keep["word"], len(senses), len(MEANING_COLS)))
            for entry in ranked:
                report.append("     kept %s row %d" % (entry["sheet"], entry["row"]))
            report.append("")
            continue

        sheets = {e["sheet"] for e in entries}
        if len(sheets) == 1:
            same_sheet += 1
        else:
            cross_sheet += 1

        keep_ws = wb[keep["sheet"]]
        existing = [meaning_key(str(keep_ws.cell(row=keep["row"], column=mc).value or "").strip())
                    for mc, _ in MEANING_COLS]
        gained = [s for s in senses if meaning_key(s[0]) not in existing]
        if gained:
            moved += len(gained)
        merges.append((keep["sheet"], keep["row"], senses))

        report.append("%s (%s)" % (keep["word"], "same sheet" if len(sheets) == 1 else "across sheets"))
        report.append("   keep   %s row %d  (%d fields, %d chars)"
                      % (keep["sheet"], keep["row"], keep["filled"], keep["length"]))
        for g in gained:
            report.append("   + kept sense from a duplicate: %s" % g[0])
        for d in drop:
            report.append("   remove %s row %d  (%d fields, %d chars)"
                          % (d["sheet"], d["row"], d["filled"], d["length"]))
            to_delete[d["sheet"]].append(d["row"])
        report.append("")

    removed = sum(len(r) for r in to_delete.values())
    if apply_changes:
        # write the merged senses before deleting, while row numbers still hold
        for sheet, row, senses in merges:
            ws = wb[sheet]
            for (mc, ec), sense in zip(MEANING_COLS, senses):
                ws.cell(row=row, column=mc).value = sense[0]
                ws.cell(row=row, column=ec).value = sense[1] or None
            for mc, ec in MEANING_COLS[len(senses):]:
                ws.cell(row=row, column=mc).value = None
                ws.cell(row=row, column=ec).value = None
        for sheet, rows in to_delete.items():
            ws = wb[sheet]
            for row in sorted(rows, reverse=True):  # bottom-up keeps indexes valid
                ws.delete_rows(row, 1)
        wb.save(FILE)

    with open("dedupe_report.txt", "w", encoding="utf-8") as fh:
        fh.write("\n".join(report))

    print("duplicate words merged, same sheet:    %d" % same_sheet)
    print("duplicate words merged, across sheets: %d" % cross_sheet)
    print("senses rescued from removed rows:      %d" % moved)
    print("words left alone (too many senses):    %d" % skipped)
    print("rows removed:                          %d" % removed)
    print("details written to dedupe_report.txt")
    print("MODE: %s" % ("WRITTEN" if apply_changes else "dry run, nothing saved"))


if __name__ == "__main__":
    main("--apply" in sys.argv)
